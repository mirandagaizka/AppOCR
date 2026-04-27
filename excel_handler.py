"""
Genera el Excel en memoria a partir de los registros de la base de datos.
No mantiene ningún archivo en disco — evita pérdidas en filesystems efímeros.
"""

import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

HEADERS = [
    "ID", "Timestamp", "N° Factura", "Fecha Factura", "Fecha Vencimiento",
    "Emisor", "NIF Emisor", "Dirección Emisor",
    "Receptor", "NIF Receptor", "Dirección Receptor",
    "Concepto", "Base Imponible", "% IVA", "Cuota IVA", "Total",
    "Forma Pago", "IBAN", "Moneda",
]

# Orden de columnas → campos del dict de factura
FIELDS = [
    "id", "timestamp", "numero_factura", "fecha_factura", "fecha_vencimiento",
    "emisor_nombre", "emisor_nif", "emisor_direccion",
    "receptor_nombre", "receptor_nif", "receptor_direccion",
    "concepto", "base_imponible", "porcentaje_iva", "cuota_iva", "total",
    "forma_pago", "iban", "moneda",
]

COL_WIDTHS = [7, 20, 15, 13, 16, 28, 14, 32, 28, 14, 32, 35, 14, 7, 11, 14, 16, 26, 8]
NUMERIC_FIELDS = {"base_imponible", "cuota_iva", "total"}

BLUE_DARK  = "1F4E79"
BLUE_LIGHT = "D6E4F0"
BORDER_CLR = "AEC6CF"
THIN   = Side(style="thin", color=BORDER_CLR)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def generate_excel_bytes(invoices: list[dict]) -> bytes:
    """Genera el Excel completo en memoria y devuelve los bytes."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Facturas"
    ws.freeze_panes = "A2"

    # Cabecera
    hdr_font  = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
    hdr_fill  = PatternFill(start_color=BLUE_DARK, end_color=BLUE_DARK, fill_type="solid")
    hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col, (header, width) in enumerate(zip(HEADERS, COL_WIDTHS), 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font  = hdr_font
        cell.fill  = hdr_fill
        cell.alignment = hdr_align
        cell.border = BORDER
        ws.column_dimensions[cell.column_letter].width = width

    ws.row_dimensions[1].height = 28

    # Filas de datos
    for row_idx, inv in enumerate(invoices, 2):
        bg = BLUE_LIGHT if row_idx % 2 == 0 else "FFFFFF"
        row_fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")

        for col, field in enumerate(FIELDS, 1):
            val = inv.get(field)
            # Convertir string numérico si viene de SQLite
            if field in NUMERIC_FIELDS and isinstance(val, str):
                try:
                    val = float(val)
                except (ValueError, TypeError):
                    val = None

            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.fill   = row_fill
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if field in NUMERIC_FIELDS and val is not None:
                cell.number_format = "#,##0.00"

        ws.row_dimensions[row_idx].height = 18

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
